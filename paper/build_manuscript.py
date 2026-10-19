from __future__ import annotations

import json
import re
import shutil
from pathlib import Path

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch, FancyArrowPatch
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
OUT = ROOT / 'build'
OUT.mkdir(exist_ok=True)
SOURCE = ROOT / 'agents_with_keys_submission.md'


def parse(text):
    text=re.sub(r'(!\[.*?\]\(.*?\))\n\n(Figure \d+\.[^\n]+)',r'\2\n\n\1',text)
    lines=text.splitlines(); blocks=[]; i=0
    while i<len(lines):
        line=lines[i].strip()
        if not line:
            i+=1; continue
        if line.startswith('```'):
            lang=line[3:]; code=[];i+=1
            while i<len(lines) and not lines[i].strip().startswith('```'):
                code.append(lines[i]);i+=1
            blocks.append(('code',lang,'\n'.join(code)));i+=1
        elif line.startswith('#'):
            m=re.match(r'^(#+)\s+(.*)',line)
            blocks.append(('heading',len(m[1]),m[2]));i+=1
        elif line.startswith('|'):
            rows=[]
            while i<len(lines) and lines[i].strip().startswith('|'):
                row=[x.strip() for x in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-+:?',x.replace(' ','')) for x in row):rows.append(row)
                i+=1
            blocks.append(('table',rows))
        elif line.startswith('!['):
            m=re.match(r'!\[(.*?)\]\((.*?)\)',line)
            blocks.append(('image',m[1],m[2]));i+=1
        else:
            p=[line];i+=1
            while i<len(lines) and lines[i].strip() and not lines[i].startswith(('#','|','![')):
                p.append(lines[i].strip());i+=1
            blocks.append(('paragraph',' '.join(p)))
    return blocks


def add_inline(p,text,fontsize=None):
    for token in re.split(r'(\*\*.+?\*\*|`[^`]+`)',text):
        if not token:continue
        r=p.add_run(token[2:-2] if token.startswith('**') else token[1:-1] if token.startswith('`') else token)
        if token.startswith('**'):r.bold=True
        if token.startswith('`'):
            r.font.name='Courier New';r.font.size=Pt((fontsize or 11)-1.5)
        elif fontsize:r.font.size=Pt(fontsize)
        r.font.color.rgb=RGBColor(0,0,0)


def make_docx(blocks):
    doc=Document();sec=doc.sections[0]
    sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=Inches(.72);sec.bottom_margin=Inches(.72)
    sec.left_margin=Inches(.82);sec.right_margin=Inches(.82)
    sec.footer_distance=Inches(.30)
    for name in ('Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption'):
        s=doc.styles[name];s.font.name='Times New Roman';s.font.color.rgb=RGBColor(0,0,0)
        s._element.get_or_add_rPr().rFonts.set(qn('w:ascii'),'Times New Roman')
        s._element.rPr.rFonts.set(qn('w:hAnsi'),'Times New Roman')
        for attr in list(s._element.rPr.rFonts.attrib):
            if 'theme' in attr.lower():del s._element.rPr.rFonts.attrib[attr]
        if s._element.pPr is not None:
            for border in list(s._element.pPr.findall(qn('w:pBdr'))):s._element.pPr.remove(border)
    for color in doc.styles.element.iter(qn('w:color')):
        for attr in list(color.attrib):
            if 'theme' in attr.lower():del color.attrib[attr]
    normal=doc.styles['Normal'];normal.font.size=Pt(11)
    normal.paragraph_format.line_spacing=1.04;normal.paragraph_format.space_after=Pt(5)
    for name,size,before,after in [('Title',18,0,9),('Heading 1',12,10,4),('Heading 2',11,7,3),('Caption',10,5,4)]:
        s=doc.styles[name];s.font.size=Pt(size)
        s.paragraph_format.space_before=Pt(before);s.paragraph_format.space_after=Pt(after)
        s.paragraph_format.keep_with_next=name!='Caption'
    doc.styles['Title'].font.bold=True
    doc.styles['Heading 1'].font.bold=True
    doc.styles['Heading 2'].font.bold=True
    doc.styles['Caption'].font.italic=True
    doc.core_properties.title=blocks[0][2]
    doc.core_properties.author='Osita Sean Keluo-Udeke'
    doc.core_properties.subject='Research manuscript on authority boundaries for AI enabled continuous integration'
    doc.core_properties.keywords='agentic engineering; continuous integration; least privilege; prompt injection'
    fp=sec.footer.paragraphs[0];fp.alignment=WD_ALIGN_PARAGRAPH.CENTER
    run=fp.add_run();run.font.size=Pt(9)
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');run._r.addnext(fld)
    refs=False;caption=None;abstract=False
    for b in blocks:
        if b[0]=='heading':
            level,title=b[1],b[2]
            if level==1:
                p=doc.add_paragraph(title,style='Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            else:
                p=doc.add_paragraph(title,style='Heading 1' if level==2 else 'Heading 2')
                refs=title=='References';abstract=title=='Abstract'
        elif b[0]=='paragraph':
            text=b[1]
            if re.match(r'^(Table|Figure) \d+\.',text):caption=text;continue
            p=doc.add_paragraph()
            if text=='Osita Sean Keluo-Udeke':
                p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(9)
            if refs:
                p.paragraph_format.left_indent=Inches(.25)
                p.paragraph_format.first_line_indent=Inches(-.25)
                p.paragraph_format.space_after=Pt(4)
                p.paragraph_format.line_spacing=1.0
            add_inline(p,text,9.5 if refs else 10.5 if abstract else 11)
            if text.startswith('Keywords:'):abstract=False
        elif b[0]=='table':
            if caption:
                p=doc.add_paragraph(caption,style='Caption');p.paragraph_format.keep_with_next=True;caption=None
            rows=b[1];n=len(rows[0]);table=doc.add_table(rows=1,cols=n)
            table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
            widths={2:[4.5,2.3],3:[1.3,2.5,3.0],4:[3.15,.55,.95,2.15],6:[1.25,.9,.9,1.05,1.15,1.55]}[n]
            for col,width in zip(table.columns,widths):col.width=Inches(width)
            for j,row in enumerate(rows):
                cells=table.rows[0].cells if j==0 else table.add_row().cells
                for k,(cell,value) in enumerate(zip(cells,row)):
                    cell.width=Inches(widths[k]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(0)
                    p.paragraph_format.line_spacing=1.03
                    if n==4 and k>0:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                    add_inline(p,value,8.5 if n==6 else 10)
                    if j==0:
                        for r in p.runs:r.bold=True;r.font.color.rgb=RGBColor(255,255,255)
                    tcPr=cell._tc.get_or_add_tcPr()
                    shd=OxmlElement('w:shd');shd.set(qn('w:fill'),'333333' if j==0 else 'F5F5F5' if j%2==0 else 'FFFFFF');tcPr.append(shd)
                    margins=OxmlElement('w:tcMar')
                    for side in ('top','left','bottom','right'):
                        m=OxmlElement('w:'+side);m.set(qn('w:w'),'100');m.set(qn('w:type'),'dxa');margins.append(m)
                    tcPr.append(margins)
                    borders=OxmlElement('w:tcBorders')
                    for side in ('top','left','bottom','right'):
                        e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
                    tcPr.append(borders)
                trPr=table.rows[j]._tr.get_or_add_trPr();trPr.append(OxmlElement('w:cantSplit'))
                if j==0:trPr.append(OxmlElement('w:tblHeader'))
            p=doc.add_paragraph();p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=Pt(2);p.add_run().font.size=Pt(2)
        elif b[0]=='code':
            p=doc.add_paragraph();p.paragraph_format.space_after=Pt(6)
            p.paragraph_format.keep_together=True
            r=p.add_run(b[2]);r.font.name='Courier New';r.font.size=Pt(9)
        elif b[0]=='image':
            p=doc.add_paragraph();p.paragraph_format.keep_with_next=True
            inline=p.add_run().add_picture(str(OUT/b[2]),width=Inches(6.80))
            inline._inline.docPr.set('descr',b[1])
            p=doc.add_paragraph((caption or 'Figure 1. Proposed trust boundaries'),style='Caption')
            p.paragraph_format.keep_with_next=False;caption=None
    dest=OUT/'Agents_With_Keys_Submission.docx';doc.save(dest)
    return dest


def tex_inline(text,citations=True):
    result=[]
    for token in re.split(r'(\*\*.+?\*\*|`[^`]+`|https?://\S+|\[\d+\])',text):
        if not token:continue
        if token.startswith('**'):
            result.append(r'\textbf{'+tex_inline(token[2:-2],citations)+'}')
        elif token.startswith('`'):
            code=token[1:-1]
            if len(code)>38:result.append(r'\nolinkurl{'+code+'}')
            else:result.append(r'\texttt{'+tex_escape(code)+'}')
        elif token.startswith(('https://','http://')):
            result.append(r'\url{'+token+'}')
        elif re.fullmatch(r'\[\d+\]',token) and citations:
            result.append(r'\cite{r'+token[1:-1]+'}')
        else:result.append(tex_escape(token))
    return ''.join(result)


def tex_escape(text):
    mapping={'\\':r'\textbackslash{}','&':r'\&','%':r'\%','$':r'\$','#':r'\#',
             '_':r'\_','{':r'\{','}':r'\}','~':r'\textasciitilde{}','^':r'\textasciicircum{}'}
    return ''.join(mapping.get(c,c) for c in text)


def make_tex(blocks):
    title=blocks[0][2]
    output=[r'\documentclass[conference,letterpaper]{IEEEtran}',r'\usepackage[T1]{fontenc}',
            r'\usepackage{times}',r'\usepackage{graphicx,array,balance}',r'\usepackage[table]{xcolor}',r'\usepackage{xurl}',
            r'\usepackage[hidelinks,breaklinks]{hyperref}',r'\urlstyle{same}',
            r'\renewcommand{\thetable}{\arabic{table}}',r'\setlength{\emergencystretch}{1em}',
            r'\raggedbottom',r'\begin{document}',r'\title{'+tex_escape(title)+'}',
            r'\author{\IEEEauthorblockN{Osita Sean Keluo-Udeke}}',r'\maketitle']
    refs=False;abstract=False;caption=None;need_close=False
    for b in blocks[2:]:
        if b[0]=='heading':
            level,title=b[1],b[2]
            if abstract:output.append(r'\end{abstract}');abstract=False
            if title=='Abstract':output.append(r'\begin{abstract}');abstract=True
            elif title=='References':output.append(r'\begin{thebibliography}{17}');refs=True
            elif title=='Acknowledgments':output.extend([r'\balance',r'\section*{Acknowledgments}'])
            else:
                title=re.sub(r'^\d+(?:\.\d+)?\s+','',title)
                output.append(('\\section{' if level==2 else '\\subsection{')+tex_escape(title)+'}')
        elif b[0]=='paragraph':
            text=b[1]
            if text.startswith('Keywords:'):
                if abstract:output.append(r'\end{abstract}');abstract=False
                output.append(r'\begin{IEEEkeywords}'+tex_inline(text.removeprefix('Keywords: '))+r'\end{IEEEkeywords}')
            elif re.match(r'^(Table|Figure) \d+\.',text):caption=re.sub(r'^(Table|Figure) \d+\.\s*','',text)
            elif refs:
                m=re.match(r'\[(\d+)\]\s+(.*)',text)
                output.append(r'\bibitem{r'+m[1]+'} '+tex_inline(m[2],False))
            else:output.append(tex_inline(text)+'\n')
        elif b[0]=='table':
            rows=b[1];n=len(rows[0]);widths={2:[.63,.29],3:[.18,.34,.38],4:[.44,.08,.14,.19],6:[.17,.13,.13,.16,.14,.18]}[n]
            env='table*' if n==6 else 'table'
            cols='|'.join(r'>{\raggedright\arraybackslash}p{'+str(w)+r'\linewidth}' for w in widths)
            output.extend(['\\begin{'+env+'}[t]',r'\caption{'+tex_inline(caption or '')+'}',r'\centering\footnotesize',
                           r'\setlength{\tabcolsep}{3pt}\renewcommand{\arraystretch}{1.12}',
                           r'\arrayrulecolor[HTML]{D9D9D9}',r'\begin{tabular}{|'+cols+'|}',r'\hline'])
            for i,row in enumerate(rows):
                values=[tex_inline(v) for v in row]
                if i==0:values=[r'\textbf{'+v+'}' for v in values]
                output.append(' & '.join(values)+r' \\ \hline')
            output.extend([r'\end{tabular}','\\end{'+env+'}']);caption=None
        elif b[0]=='code':
            output.extend([r'\begingroup\scriptsize',r'\begin{verbatim}',b[2],r'\end{verbatim}',r'\endgroup'])
        elif b[0]=='image':
            output.extend([r'\begin{figure*}[t]',r'\centering',r'\includegraphics[width=\textwidth]{figure1.pdf}',
                           r'\caption{'+tex_inline(caption or '')+'}',r'\end{figure*}']);caption=None
    output.extend([r'\end{thebibliography}',r'\end{document}'])
    (OUT/'manuscript.tex').write_text('\n\n'.join(output)+'\n')


def main():
    import argparse
    from figure1 import render
    global OUT,SOURCE
    p=argparse.ArgumentParser()
    p.add_argument('--source',type=Path,default=SOURCE)
    p.add_argument('--output',type=Path,default=OUT)
    p.add_argument('--artifact-commit')
    args=p.parse_args();OUT=args.output.resolve();OUT.mkdir(parents=True,exist_ok=True)
    SOURCE=args.source.resolve();text=SOURCE.read_text()
    if args.artifact_commit:
        if not re.fullmatch(r'[0-9a-f]{40}',args.artifact_commit):raise ValueError('invalid artifact commit')
        text=text.replace('The artifact is version 0.1.0 (tag `v0.1.0`)',
                          'The artifact is version 0.1.0 (tag `v0.1.0`), commit `'+args.artifact_commit+'`')
    render(OUT);blocks=parse(text);make_docx(blocks);make_tex(blocks)
    shutil.copyfile(ROOT/'IEEEtran.cls',OUT/'IEEEtran.cls')
    print(json.dumps({'docx':str(OUT/'Agents_With_Keys_Submission.docx'),'source_words':len(text.split())}))


if __name__=='__main__':main()
